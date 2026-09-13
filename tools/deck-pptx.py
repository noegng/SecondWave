#!/usr/bin/env python3
"""tools/deck-pptx.py — SLIDES.html → SecondWave.pptx

    npm run deck        (nécessite python-pptx : pip install python-pptx)

Le deck HTML reste la référence visuelle ; celui-ci existe parce qu'un .pptx
s'édite à la main et s'importe dans Google Slides tel quel.

Le contenu est dupliqué ici plutôt que scrapé du HTML : les deux supports ne
mettent pas les mêmes choses au même endroit, et un parseur HTML fragile
serait plus coûteux à maintenir que quinze lignes de texte.

Palette et esprit repris du front (apps/web/public/styles.css) : fond noir,
texte blanc, les rouges pour ce qui coince et les verts pour ce qu'on apporte.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── palette ───────────────────────────────────────────────────────────────
BG     = RGBColor(0x0B, 0x0B, 0x0D)
INK    = RGBColor(0xF1, 0xF1, 0xF3)
STEEL  = RGBColor(0xC4, 0xC4, 0xCA)
MUTED  = RGBColor(0x9A, 0x9A, 0xA1)
DIM    = RGBColor(0x6F, 0x6F, 0x77)
LINE   = RGBColor(0x2B, 0x2B, 0x30)
BAD    = RGBColor(0xFF, 0x5A, 0x52)
GOOD   = RGBColor(0x35, 0xD0, 0x7F)
WARN   = RGBColor(0xF5, 0xC5, 0x42)

FONT = 'Helvetica Neue'
MONO = 'Menlo'

W, H = Inches(13.333), Inches(7.5)
L    = Inches(0.95)                      # marge gauche
CW   = W - Inches(1.9)                   # largeur utile

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ── briques ───────────────────────────────────────────────────────────────
def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def vault_motif(s):
    """Les arêtes du coffre en perspective, comme sur le front."""
    def rect(x, y, w, h):
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        r.fill.background(); r.line.color.rgb = LINE; r.line.width = Pt(0.75)
        r.shadow.inherit = False
        return r

    def line(x1, y1, x2, y2):
        c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = LINE; c.line.width = Pt(0.75)
        return c

    rect(0.45, 0.35, 12.45, 6.8)
    rect(1.15, 0.95, 11.05, 5.6)
    line(0.45, 0.35, 1.15, 0.95)
    line(12.90, 0.35, 12.20, 0.95)
    line(0.45, 7.15, 1.15, 6.55)
    line(12.90, 7.15, 12.20, 6.55)


def box(s, x, y, w, h):
    # Une zone de texte qui déborde du cadre se fait rogner à l'export PDF et
    # à l'import Google Slides. On la borne plutôt que de compter les pouces
    # à la main sur chaque diapositive.
    h = min(h, H - y - Inches(0.25))
    w = min(w, W - x - Inches(0.25))
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def runs(p, text, size, color, font=FONT, bold=False, space_after=Pt(0)):
    """Écrit `text` dans le paragraphe `p`. `**gras**` et `` `mono` `` reconnus."""
    p.space_after = space_after
    import re
    for frag in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not frag:
            continue
        r = p.add_run()
        if frag.startswith('**'):
            r.text = frag[2:-2]; r.font.bold = True; r.font.color.rgb = INK
            r.font.name = font
        elif frag.startswith('`'):
            r.text = frag[1:-1]; r.font.name = MONO; r.font.color.rgb = STEEL
        else:
            r.text = frag; r.font.bold = bold; r.font.color.rgb = color
            r.font.name = font
        r.font.size = size
    return p


def kicker(s, text, y=0.75):
    tf = box(s, L, Inches(y), CW, Inches(0.35))
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text.upper()
    r.font.size = Pt(11); r.font.name = MONO; r.font.color.rgb = DIM
    f = r.font._rPr
    f.set('spc', '260')                  # letter-spacing, absent de l'API
    return tf


def title(s, text, y=1.35, size=34, width=None, color=INK):
    tf = box(s, L, Inches(y), width or Inches(10.4), Inches(1.8))
    p = tf.paragraphs[0]; p.line_spacing = 1.04
    runs(p, text, Pt(size), color, bold=True)
    return tf


def lede(s, text, y, size=17, width=None):
    tf = box(s, L, Inches(y), width or Inches(9.6), Inches(1.4))
    p = tf.paragraphs[0]; p.line_spacing = 1.35
    runs(p, text, Pt(size), STEEL)
    return tf


def bullets(s, items, y, size=15, width=None, gap=0.06):
    """items : (texte, ton) avec ton ∈ {None, 'bad', 'good'}."""
    tf = box(s, L, Inches(y), width or Inches(10.6), Inches(3.6))
    first = True
    for text, tone in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.32
        p.space_after = Pt(gap * 72)
        d = p.add_run(); d.text = '—   '
        d.font.size = Pt(size); d.font.name = FONT
        d.font.color.rgb = {'bad': BAD, 'good': GOOD}.get(tone, DIM)
        d.font.bold = tone is not None
        runs(p, text, Pt(size), STEEL)
    return tf


def card(s, x, y, w, h, number, name, funcs):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = RGBColor(0x13, 0x13, 0x15)
    r.line.color.rgb = LINE; r.line.width = Pt(0.75)
    r.shadow.inherit = False
    r.adjustments[0] = 0.06

    tf = box(s, Inches(x + 0.28), Inches(y + 0.22), Inches(w - 0.5), Inches(h - 0.4))
    p = tf.paragraphs[0]; p.space_after = Pt(2)
    a = p.add_run(); a.text = number
    a.font.size = Pt(10); a.font.name = MONO; a.font.color.rgb = STEEL; a.font.bold = True
    p2 = tf.add_paragraph(); p2.space_after = Pt(5)
    b = p2.add_run(); b.text = name
    b.font.size = Pt(13); b.font.name = FONT; b.font.color.rgb = INK; b.font.bold = True
    for f in funcs:
        p3 = tf.add_paragraph(); p3.space_after = Pt(0); p3.line_spacing = 1.25
        c = p3.add_run(); c.text = f
        c.font.size = Pt(9.5); c.font.name = MONO; c.font.color.rgb = MUTED


def num(s, n):
    tf = box(s, W - Inches(1.6), H - Inches(0.75), Inches(0.9), Inches(0.35))
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = f'{n:02d}'
    r.font.size = Pt(10); r.font.name = MONO; r.font.color.rgb = DIM


# ══ 1 · titre ═════════════════════════════════════════════════════════════
s = slide(); vault_motif(s)
kicker(s, 'XRPL · Track 2 — Lending', 1.5)
tf = box(s, L, Inches(2.0), Inches(11), Inches(1.5))
p = tf.paragraphs[0]
r = p.add_run(); r.text = 'SecondWave'
r.font.size = Pt(66); r.font.bold = True; r.font.color.rgb = INK; r.font.name = FONT
lede(s, 'An exit market for locked vault shares — and a risk rating that says '
        'whether the discount is an opportunity or a trap.', 3.5, 17, Inches(8.8))
tf = box(s, L, Inches(5.0), Inches(9), Inches(1.0))
for i, (who, what) in enumerate([('Hugo', ' — vaults, loans, settlement'),
                                 ('Noé', ' — risk analyst, wallet')]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(3)
    a = p.add_run(); a.text = who
    a.font.size = Pt(15); a.font.color.rgb = STEEL; a.font.name = FONT; a.font.bold = True
    b = p.add_run(); b.text = what
    b.font.size = Pt(15); b.font.color.rgb = DIM; b.font.name = FONT
tf = box(s, L, Inches(6.1), Inches(9), Inches(0.4))
r = tf.paragraphs[0].add_run(); r.text = 'XRPL PUBLIC DEVNET · SEPTEMBER 2026'
r.font.size = Pt(10); r.font.name = MONO; r.font.color.rgb = DIM
num(s, 1)

# ══ 2 · le problème ═══════════════════════════════════════════════════════
s = slide()
kicker(s, 'The observation')
title(s, 'A closed-ended vault locks your capital.\nThat is the product, not a bug.', 1.3, 33)
lede(s, 'During the whole Investment phase, `VaultWithdraw` answers `tecTOO_SOON`. '
        'The depositor cannot leave. Full stop.', 3.0, 16, Inches(10.2))
bullets(s, [
    ('But the **risk degrades while the price does not move**. `AssetsTotal` only changes '
     'when the broker declares an impairment — and nothing forces them to.', 'bad'),
    ('A loan past its due date and its grace period leaves the share value perfectly intact. '
     '**It is invisible from the price.**', 'bad'),
    ('We swept the network: **22 of the 40 largest vaults** carry an undeclared defaulted loan. '
     'Not ours. In the wild.', None),
], 4.15, 15, Inches(11.1))
num(s, 2)

# ══ 3 · la solution ═══════════════════════════════════════════════════════
s = slide()
kicker(s, 'What we built')
title(s, 'Vault shares are an MPT. They transfer.\nThe exit already existed — it just had no tooling.', 1.3, 31)
bullets(s, [
    ('**Sell a locked position** before the redemption date, through an all-or-nothing atomic swap.', 'good'),
    ('**An analyst next to every offer** — not just “risky”, but distress discount versus '
     'liquidity discount, with the reason.', 'good'),
    ('**A durable order book.** The seller signs once and keeps living; the buyer commits with a '
     'deadline; either side can withdraw in one click, on-chain.', 'good'),
], 3.25, 16, Inches(11.1), gap=0.1)
tf = box(s, L, Inches(5.95), Inches(10.6), Inches(0.9))
p = tf.paragraphs[0]; p.line_spacing = 1.35
runs(p, 'The buyer’s real question is not what price — it is **why is the seller cutting it**. '
        'That is what the rating answers.', Pt(14), MUTED)
num(s, 3)

# ══ 4 · les XLS ═══════════════════════════════════════════════════════════
s = slide()
kicker(s, 'What we ran on')
title(s, 'Six specifications, one cycle', 1.3, 33)
specs = [
    ('XLS-65', 'Single Asset Vault',   ['VaultCreate · VaultDeposit', 'VaultWithdraw · VaultClawback']),
    ('XLS-66', 'Lending Protocol',     ['LoanBrokerSet · LoanSet', 'LoanPay · LoanManage · LoanDelete']),
    ('XLS-56', 'Batch',                ['Batch tfAllOrNothing', 'BatchSigners · inner legs']),
    ('XLS-70', 'Credentials',          ['CredentialCreate · CredentialAccept', 'CredentialDelete']),
    ('XLS-80', 'Permissioned Domains', ['PermissionedDomainSet', 'AcceptedCredentials']),
    ('MPT · TICKETS', 'Shares & durability', ['MPTokenAuthorize · Payment(mpt)', 'TicketCreate · TicketSequence']),
]
for i, (n, name, fs) in enumerate(specs):
    card(s, 0.95 + (i % 3) * 4.0, 2.45 + (i // 3) * 1.75, 3.72, 1.5, n, name, fs)
tf = box(s, L, Inches(6.25), Inches(11), Inches(0.5))
r = tf.paragraphs[0].add_run()
r.text = 'Around 450 cases replayed on Devnet · 30 vaults created · 489 swept'
r.font.size = Pt(13); r.font.color.rgb = MUTED; r.font.name = FONT
num(s, 4)

# ══ 5 · les fixes majeurs ═════════════════════════════════════════════════
s = slide()
kicker(s, 'Major fixes we surfaced')
title(s, 'Two things a product cannot work around', 1.3, 33)
for i, (tag, head, items) in enumerate([
    ('XLS-56 · BATCH', '“Success” does not mean the trade happened', [
        'Delivery and no-op return the **same tesSUCCESS**, the same metadata, the same fee. '
        'Two Devnet transactions side by side: **rigorously identical**.',
        'The per-leg execution report XLS-56 specifies is **absent**. We rebuild the legs by '
        'hand and reconcile balances.',
        '**3 of 4 flags** hand over the shares without collecting the price.']),
    ('DEV WALLET', 'It cannot get you in — or out', [
        'It sends an XRP deposit **as a token object**. The vault refuses, and the user cannot '
        'edit the JSON. Dead end — we fixed it and opened a PR.',
        'It knows **neither Batch nor Escrow**. No wallet can sign an atomic swap, so every '
        'peer-to-peer market is forced to be **custodial**.']),
]):
    x = 0.95 + i * 6.1
    tf = box(s, Inches(x), Inches(2.45), Inches(5.6), Inches(0.3))
    r = tf.paragraphs[0].add_run(); r.text = tag
    r.font.size = Pt(10); r.font.name = MONO; r.font.color.rgb = BAD; r.font.bold = True
    tf = box(s, Inches(x), Inches(2.85), Inches(5.6), Inches(0.6))
    p = tf.paragraphs[0]; p.line_spacing = 1.15
    r = p.add_run(); r.text = head
    r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = INK; r.font.name = FONT
    tf = box(s, Inches(x), Inches(3.75), Inches(5.6), Inches(3.0))
    for k, it in enumerate(items):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.line_spacing = 1.28; p.space_after = Pt(7)
        d = p.add_run(); d.text = '—   '
        d.font.size = Pt(13); d.font.color.rgb = BAD; d.font.name = FONT; d.font.bold = True
        runs(p, it, Pt(13), STEEL)
num(s, 5)

# ══ 6 · l'escrow au porteur ═══════════════════════════════════════════════
s = slide()
kicker(s, 'What would have been great to have')
title(s, 'An escrow that does not name its recipient', 1.3, 33)
lede(s, 'Today the seller must know the buyer **before committing**: the buyer’s address sits '
        'inside what the seller signs, and an escrow demands a `Destination`. So there is no way '
        'to say “1,000 shares at this price, to whoever takes them first”.', 2.5, 15, Inches(11.1))
bullets(s, [
    ('The usual objection — anyone could claim it — **does not apply here**.', None),
    ('In a private vault, access is already bounded by a **permissioned domain** (XLS-80) and its '
     '**credentials** (XLS-70). Everyone who could answer is **already trusted**.', 'good'),
    ('The network re-checks membership at deposit, at transfer, and **again at escrow finish** — '
     'we tested it.', 'good'),
], 4.05, 14, Inches(11.1))
tf = box(s, L, Inches(6.2), Inches(11.1), Inches(0.8))
p = tf.paragraphs[0]; p.line_spacing = 1.3
runs(p, 'Naming the destination a second time adds no security. It only removes the possibility '
        'of a market. **Let an escrow point at a domain** — and private vaults get a real '
        'on-chain order book.', Pt(14), MUTED)
num(s, 6)

# ══ 7 · l'onboarding off-chain ════════════════════════════════════════════
s = slide()
kicker(s, 'The one thing still off-chain')
title(s, 'There is no way to ask to join', 1.3, 33)
lede(s, 'Everything else lives on the ledger. Onboarding does not.', 2.35, 17, Inches(10))
bullets(s, [
    ('The only transactions that exist are `CredentialCreate`, `CredentialAccept`, '
     '`CredentialDelete`. **None of them is a request.**', None),
    ('The **issuer pushes** the credential; the newcomer can only accept it. An outsider has '
     '**no on-chain way to raise their hand**.', 'bad'),
    ('So the first step of onboarding happens in a form, an email, a Discord — off-chain, for a '
     'protocol whose selling point is on-chain identity.', None),
], 3.3, 15, Inches(11.1), gap=0.09)
tf = box(s, L, Inches(6.05), Inches(11.1), Inches(0.8))
p = tf.paragraphs[0]; p.line_spacing = 1.3
runs(p, 'Concretely: a buyer who finds an offer on our book and is not a domain member sees '
        '“credential missing” — and has no button to press next.', Pt(14), MUTED)
num(s, 7)

# ══ 8 · les liens ═════════════════════════════════════════════════════════
s = slide(); vault_motif(s)
kicker(s, 'Go look', 1.5)
tf = box(s, L, Inches(2.0), Inches(11), Inches(1.2))
r = tf.paragraphs[0].add_run(); r.text = 'SecondWave'
r.font.size = Pt(52); r.font.bold = True; r.font.color.rgb = INK; r.font.name = FONT
for i, (k, v) in enumerate([('LIVE', 'second-wave-vault.vercel.app'),
                            ('CODE', 'github.com/noegng/SecondWave')]):
    y = 3.45 + i * 1.15
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, L, Inches(y), Inches(6.4), Inches(0.95))
    r.fill.background(); r.line.color.rgb = LINE; r.line.width = Pt(0.75)
    r.shadow.inherit = False; r.adjustments[0] = 0.08
    tf = box(s, Inches(1.25), Inches(y + 0.16), Inches(5.8), Inches(0.7))
    p = tf.paragraphs[0]; p.space_after = Pt(2)
    a = p.add_run(); a.text = k
    a.font.size = Pt(9); a.font.name = MONO; a.font.color.rgb = DIM
    p2 = tf.add_paragraph()
    b = p2.add_run(); b.text = v
    b.font.size = Pt(15); b.font.name = MONO; b.font.color.rgb = INK
tf = box(s, L, Inches(5.95), Inches(10.4), Inches(0.9))
p = tf.paragraphs[0]; p.line_spacing = 1.3
runs(p, 'The report, the reproduction scripts and the on-chain proofs are all in the repository — '
        '**every claim on these slides has a transaction you can open.**', Pt(14), MUTED)
num(s, 8)

out = 'SecondWave.pptx'
prs.save(out)
print(f'{out} — {len(prs.slides.__iter__.__self__._sldIdLst)} slides')
